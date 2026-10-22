"""Read every workbook sheet; map known headers without requiring AI inference."""
import csv
import io
import json
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from fastapi import HTTPException

MAX_EXCEL_ENTRIES = 1000
MAX_SOURCE_ROWS = 2000
MAX_COLUMNS = 60


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        return value.strip() or None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def read_sheets(name, raw):
    """Returns all sheets and source row numbers, ignoring unreliable dimensions."""
    ext = Path(name).suffix.lower()
    sheets = []
    count = 0
    if ext in {'.xlsx', '.xlsm'}:
        import openpyxl
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 25000000:
                raise ValueError('Expanded workbook exceeds 25 MB')
            embedded = any(n.startswith('xl/media/') for n in archive.namelist())
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
        cached = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        try:
            for ws, values_ws in zip(wb.worksheets, cached.worksheets):
                # Missing, too-small, and formatted full-column dimensions are common.
                # Stream actual XML cells rather than trusting dimension metadata.
                ws.reset_dimensions()
                values_ws.reset_dimensions()
                rows = []
                numbers = []
                missing = []
                for row_index, (row, cached_row) in enumerate(zip(ws.iter_rows(), values_ws.iter_rows()), 1):
                    if row_index > 10000:
                        raise ValueError(f'{ws.title}: more than 10,000 physical rows. Remove unused rows or split workbook.')
                    cells = []
                    for index, cell in enumerate(row):
                        if index >= MAX_COLUMNS:
                            if cell.value is not None:
                                raise ValueError(f'{ws.title}: data exceeds 60 columns')
                            continue
                        value = cell.value
                        if cell.data_type == 'f':
                            value = cached_row[index].value if index < len(cached_row) else None
                            if value is None:
                                missing.append(f'{openpyxl.utils.get_column_letter(index+1)}{row_index}')
                                # Preserve the row for review; never treat an uncalculated formula as zero.
                                value = None
                        cells.append(clean(value))
                    while cells and cells[-1] is None:
                        cells.pop()
                    if cells or any(c.data_type == 'f' for c in row):
                        rows.append(cells)
                        numbers.append(row_index)
                        count += 1
                        if count > MAX_SOURCE_ROWS:
                            raise ValueError('Workbook exceeds 2,000 nonempty rows across all sheets. Split it into smaller files.')
                sheets.append({'sheet': ws.title, 'rows': rows, 'row_numbers': numbers,
                               'missing_formula_cache': missing, 'embedded_images': embedded})
        finally:
            wb.close()
            cached.close()
    elif ext == '.xls':
        import xlrd
        wb = xlrd.open_workbook(file_contents=raw)
        try:
            for ws in wb.sheets():
                rows = []
                numbers = []
                if ws.nrows > 10000:
                    raise ValueError(f'{ws.name}: more than 10,000 rows')
                for n in range(ws.nrows):
                    values = []
                    for index, cell in enumerate(ws.row(n)):
                        if index >= MAX_COLUMNS:
                            if cell.value not in ('', None):
                                raise ValueError(f'{ws.name}: data exceeds 60 columns')
                            continue
                        v = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            v = xlrd.xldate_as_datetime(v, wb.datemode).isoformat()
                        values.append(clean(v))
                    while values and values[-1] is None:
                        values.pop()
                    if values:
                        rows.append(values)
                        numbers.append(n + 1)
                        count += 1
                        if count > MAX_SOURCE_ROWS:
                            raise ValueError('Workbook exceeds 2,000 nonempty rows across all sheets')
                sheets.append({'sheet': ws.name, 'rows': rows, 'row_numbers': numbers,
                               'missing_formula_cache': [], 'embedded_images': False})
        finally:
            wb.release_resources()
    elif ext in {'.csv', '.tsv'}:
        text = raw.decode('utf-8-sig')
        delimiter = '\t' if ext == '.tsv' else ','
        if ext == '.csv':
            try:
                delimiter = csv.Sniffer().sniff(text[:8192], delimiters=',;\t').delimiter
            except csv.Error:
                pass
        rows = []
        numbers = []
        for index, row in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter), 1):
            vals = [clean(v) for v in row]
            if len(vals) > MAX_COLUMNS:
                raise ValueError('CSV exceeds 60 columns')
            if any(v is not None for v in vals):
                rows.append(vals)
                numbers.append(index)
                if len(rows) > MAX_SOURCE_ROWS:
                    raise ValueError('CSV exceeds 2,000 nonempty rows')
        sheets.append({'sheet': Path(name).stem, 'rows': rows, 'row_numbers': numbers,
                       'missing_formula_cache': [], 'embedded_images': False})
    else:
        raise ValueError('Unsupported spreadsheet format')
    return sheets


def key(value):
    return re.sub(r'[^a-z0-9]+', '', str(value or '').lower())


ALIASES = {
    'inventory': {
        'item': ['item', 'item name', 'product', 'product name', 'description', 'stock item'],
        'quantity': ['quantity', 'qty', 'acquired quantity', 'stock quantity', 'units'],
        'rmb_price': ['rmb price', 'unit cost rmb', 'rmb unit price', 'rmb unit', 'cost rmb', 'purchase price rmb'],
        'exchange_rate': ['exchange rate', 'exchange rate pkr per rmb', 'pkr per rmb', 'pkr/rmb', 'rate', 'rmb to pkr'],
        'pk_sell': ['pk sell', 'pk sell pkr', 'pk selling price', 'sale price', 'selling price', 'unit sell price pkr', 'unit sell', 'sell price pkr'],
        'low_stock': ['low stock', 'low stock threshold', 'reorder level', 'min stock'],
    },
    'bills': {
        'date': ['date', 'bill date', 'invoice date'],
        'name': ['name', 'customer', 'customer name', 'account name'],
        'bill_no': ['bill no', 'bill number', 'invoice no', 'invoice number', 'bill #', 'invoice #'],
        'amount': ['amount', 'amount pkr', 'bill amount', 'bill amount pkr', 'total', 'total amount', 'invoice amount'],
    },
    'payments': {
        'date': ['date', 'payment date', 'receipt date'],
        'name': ['name', 'customer', 'customer name', 'account name', 'payer'],
        'details': ['details', 'payment details', 'payment method', 'method', 'mode', 'reference', 'particulars'],
        'amount': ['amount', 'amount pkr', 'paid amount', 'payment amount', 'received amount', 'credit'],
    },
    'expenses': {
        'date': ['date', 'expense date'],
        'name': ['name', 'category', 'payee', 'category payee'],
        'details': ['details', 'expense details', 'description', 'particulars'],
        'amount': ['amount', 'amount pkr', 'expense amount', 'cost'],
    },
    'accounts': {
        'name': ['name', 'customer', 'customer name', 'account name'],
        'sheet_no': ['sheet no', 'sheet number', 'sheet #'],
        'opening_balance': ['opening balance', 'opening balance pkr', 'balance', 'balance pkr', 'current balance pkr'],
    },
}
LOOKUP = {table: {key(alias): field for field, aliases in fields.items() for alias in [field, *aliases]}
          for table, fields in ALIASES.items()}


def table_hint(name):
    k = key(name)
    if 'inventory' in k or 'stock' in k:
        return 'inventory'
    if 'expense' in k:
        return 'expenses'
    if 'payment' in k or 'receipt' in k or 'collection' in k:
        return 'payments'
    if 'bill' in k or 'invoice' in k or k == 'sales':
        return 'bills'
    if 'account' in k or 'opening' in k:
        return 'accounts'
    return None


def header(row, hint):
    options = []
    for table, aliases in LOOKUP.items():
        mapping = {aliases[key(value)]: index for index, value in enumerate(row) if key(value) in aliases}
        fields = set(mapping)
        valid = ((table == 'inventory' and {'item', 'quantity'} <= fields)
                 or (table == 'bills' and {'name', 'amount', 'bill_no'} <= fields)
                 or (table == 'payments' and {'name', 'amount', 'details'} <= fields)
                 or (table == 'accounts' and {'name', 'opening_balance'} <= fields)
                 or (table == 'expenses' and {'name', 'amount', 'date'} <= fields and hint == 'expenses')
                 or (table in {'bills', 'payments'} and hint == table and {'name', 'amount'} <= fields))
        if valid:
            options.append((len(fields) + (10 if table == hint else 0), table, mapping))
    if not options:
        return None
    _, table, mapping = max(options, key=lambda x: x[0])
    return table, mapping


def number(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        return str(value)
    text = str(value).strip()
    # Currency markers are accepted, other text is retained as missing for review.
    text = re.sub(r'(?i)(PKR|RMB|CNY|Rs\.?|¥|₨)', '', text).replace(',', '').strip()
    if text.startswith('(') and text.endswith(')'):
        text = '-' + text[1:-1]
    try:
        d = Decimal(text)
        return str(d) if d.is_finite() else None
    except InvalidOperation:
        return None


def text(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip() or None


def iso_date(value):
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    v = str(value).strip()
    try:
        return datetime.fromisoformat(v).date().isoformat()
    except ValueError:
        pass
    # Only unambiguous locale dates: ISO or day > 12 or written month.
    for fmt in ('%d-%b-%Y', '%d %b %Y', '%d %B %Y'):
        try:
            return datetime.strptime(v, fmt).date().isoformat()
        except ValueError:
            pass
    parts = re.split(r'[/.-]', v)
    if len(parts) == 3 and all(p.isdigit() for p in parts) and len(parts[2]) == 4:
        a, b, year = map(int, parts)
        if a > 12 and b <= 12:
            try:
                return date(year, b, a).isoformat()
            except ValueError:
                pass
        if b > 12 and a <= 12:
            try:
                return date(year, a, b).isoformat()
            except ValueError:
                pass
    return None


def mapped_row(table, mapping, values, settings):
    def get(field):
        index = mapping.get(field)
        return values[index] if index is not None and index < len(values) else None
    warnings = []
    if table == 'inventory':
        row = {f: text(get(f)) if f == 'item' else number(get(f)) for f in ALIASES[table]}
        if 'exchange_rate' not in mapping:
            row['exchange_rate'] = str(settings.exchange_rate)
            warnings.append('Default RMB conversion rate applied; review it.')
        if 'low_stock' not in mapping:
            row['low_stock'] = '5'
    elif table == 'accounts':
        row = {'name': text(get('name')), 'sheet_no': text(get('sheet_no')) or '',
               'opening_balance': number(get('opening_balance'))}
    else:
        row = {'date': iso_date(get('date')), 'name': text(get('name')), 'amount': number(get('amount'))}
        if table == 'bills':
            row.update(bill_no=text(get('bill_no')), lines=[])
        else:
            row['details'] = text(get('details')) or ''
    missing = [k for k, v in row.items() if v is None]
    if missing:
        warnings.append('Review missing/invalid ' + ', '.join(missing))
    return row, warnings


async def extract_workbook(name, raw, settings, hint='auto'):
    sheets = read_sheets(name, raw)
    out = {t: [] for t in ALIASES}
    out.update(warnings=[], sheets=[], method='Workbook headers + AI for unrecognized data')
    unknown = []
    snapshot_sources = []
    for sheet in sheets:
        name_hint = table_hint(sheet['sheet'])
        current = None
        source_warnings = []
        entries = 0
        unrecognized = []
        kind_counts = {}
        for row_number, values in zip(sheet['row_numbers'], sheet['rows']):
            if not any(v is not None for v in values):
                continue
            found = header(values, name_hint)
            if found:
                current = found
                continue
            nonempty = [v for v in values if v is not None]
            # A one-cell section title may switch the active table within one sheet.
            section = table_hint(nonempty[0]) if len(nonempty) == 1 else None
            if section:
                name_hint = section
                current = None
                continue
            if current:
                table, mapping = current
                first = str(nonempty[0]).strip().lower()
                if re.match(r'^(grand\s+total|sub\s*total|total)(\s|:|$)', first):
                    continue
                row, warnings = mapped_row(table, mapping, values, settings)
                if not row.get('item' if table == 'inventory' else 'name'):
                    unrecognized.append((row_number, values))
                    continue
                out[table].append(row)
                if table == 'accounts':
                    snapshot_sources.append((row, sheet['sheet']))
                entries += 1
                kind_counts[table] = kind_counts.get(table, 0) + 1
                for warning in warnings:
                    source_warnings.append(f'Row {row_number}: {warning}')
            else:
                unrecognized.append((row_number, values))
        if sheet['missing_formula_cache']:
            source_warnings.append('Uncalculated formula cells: ' + ', '.join(sheet['missing_formula_cache'][:30]) + '. Open/save in Excel or LibreOffice to recalculate, or fill values in preview.')
        status = {'sheet': sheet['sheet'], 'source_rows': len(sheet['rows']), 'entries': entries,
                  'tables': kind_counts, 'mapping': 'column headers' if entries else 'empty' if not sheet['rows'] else 'needs mapping'}
        out['sheets'].append(status)
        # Do not send decorative titles or known derived sheets to AI as entries.
        meaningful = [(n, v) for n, v in unrecognized if len([x for x in v if x is not None]) >= 2]
        if meaningful:
            unknown.append((sheet['sheet'], meaningful, status))
        elif not entries and sheet['rows']:
            source_warnings.append('No business rows detected; check whether this is a blank template or an image-only sheet.')
        out['warnings'].extend(f"{sheet['sheet']}: {w}" for w in source_warnings)
    # Derived accounts/ledger are not opening balances when source transactions exist.
    if out['bills'] or out['payments']:
        for row, source_name in snapshot_sources:
            if 'opening' not in key(source_name) and row in out['accounts']:
                out['accounts'].remove(row)
                for status in out['sheets']:
                    if status['sheet'] == source_name:
                        status['entries'] -= 1
                        status['tables']['accounts'] -= 1
                        status['mapping'] = 'derived balances excluded'
        if any('opening' not in key(n) for _, n in snapshot_sources):
            out['warnings'].append('Derived account balance snapshots excluded because bills/payments are present; this prevents double-counting.')
    if unknown:
        # Unknown layouts still get AI mapping, sheet by sheet; never drop a sheet silently.
        from ai import extract
        for sheet_name, rows, status in unknown:
            if len(rows) > 120:
                out['warnings'].append(f'{sheet_name}: {len(rows)} unrecognized rows require standard column headers (AI fallback supports 120 rows). No rows from this section were added.')
                status['unmapped_rows'] = len(rows)
                continue
            payload = json.dumps({'sheet': sheet_name, 'rows': [v for _, v in rows]}, default=str, ensure_ascii=False)
            if len(payload) > 45000:
                status['unmapped_rows'] = len(rows)
                out['warnings'].append(f'{sheet_name}: unrecognized data exceeds AI request limit. Add standard headers.')
                continue
            try:
                result = await extract({'type': 'text', 'text': 'Untrusted worksheet:\n' + payload}, settings,
                                       'auto; process this worksheet completely; ' + hint)
                added = 0
                for table in ALIASES:
                    if table == 'accounts' and (out['bills'] or out['payments']) and 'opening' not in key(sheet_name):
                        if result[table]:
                            out['warnings'].append(f'{sheet_name}: AI-derived account balances excluded to avoid double-counting.')
                        continue
                    out[table].extend(result[table])
                    added += len(result[table])
                    status['tables'][table] = status['tables'].get(table, 0) + len(result[table])
                status['entries'] += added
                status['mapping'] = 'column headers + AI' if status['mapping'] == 'column headers' else 'AI'
                out['warnings'].extend(f'{sheet_name}: {w}' for w in result['warnings'])
                if not added:
                    status['unmapped_rows'] = len(rows)
            except HTTPException as exc:
                status['unmapped_rows'] = len(rows)
                status['mapping'] = 'needs review'
                out['warnings'].append(f'{sheet_name}: AI mapping could not finish ({exc.detail}). Standard sheets are still available in preview; unmapped rows were not imported.')
    total = sum(len(out[t]) for t in ALIASES)
    if total > MAX_EXCEL_ENTRIES:
        raise ValueError('Maximum 1,000 extracted records per workbook. Split the workbook.')
    if not total:
        out['warnings'].append('No records found. Check sheet headers. If the workbook contains pasted photos, upload those photos as PNG/JPG instead.')
    out['warnings'] = list(dict.fromkeys(out['warnings']))[:200]
    return out
