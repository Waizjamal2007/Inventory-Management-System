import io, json, sys
from decimal import Decimal
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import pytest
from fastapi.testclient import TestClient
from models import State, normalize, reports
from ai import read_upload

@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setenv('LOCAL_DB_PATH',str(tmp_path/'test.sqlite3'))
    monkeypatch.delenv('UPSTASH_REDIS_REST_URL',raising=False)
    monkeypatch.delenv('UPSTASH_REDIS_REST_TOKEN',raising=False)
    monkeypatch.delenv('VERCEL',raising=False)
    import server
    return TestClient(server.app)

def login(c):
    r=c.post('/api/login',json={'username':'admin@123','password':'admin@123'})
    assert r.status_code==200,r.text
    c.headers['x-csrf-token']=r.json()['csrf']

def sample():
    return State.model_validate({'inventory':[{'id':'i','item':'Shirt','quantity':'10','rmb_price':'10','exchange_rate':'40','pk_sell':'700'}], 'bills':[{'id':'b','date':'2026-10-22','name':'Waiz','bill_no':'03','amount':'0','lines':[{'item_id':'i','quantity':'2','price':'700'}]}], 'payments':[{'id':'p','date':'2026-10-22','name':'waiz','details':'Cash','amount':'500'}], 'expenses':[{'id':'e','date':'2026-10-22','name':'Transport','amount':'100'}]})

def test_stock_profit_and_account_chain():
    s=normalize(sample());r=reports(s)
    assert r['inventory'][0]['remaining']==8
    assert r['totals']['revenue']==1400
    assert r['totals']['net_profit']==500
    assert r['totals']['receivables']==900
    assert r['totals']['stock_value']==3200
    assert r['accounts'][0]['balance']==900

def test_cost_snapshot_and_overselling():
    old=normalize(sample());new=old.model_copy(deep=True)
    new.inventory[0].rmb_price=Decimal('99')
    normalize(new,old)
    assert new.bills[0].lines[0].unit_cost==400
    new.bills[0].lines[0].quantity=Decimal(100)
    with pytest.raises(ValueError,match='Insufficient'):normalize(new,old)

def test_uncosted_sales_not_reported_as_profit():
    s=sample();s.bills[0].lines=[];s.bills[0].amount=Decimal(20000)
    r=reports(normalize(s))
    assert r['totals']['uncosted_sales']==20000
    assert r['totals']['net_profit']==-100

def test_session_csrf_and_persistence(client):
    assert client.get('/api/state').status_code==401
    assert client.post('/api/login',json={'username':'wrong','password':'wrong'}).status_code==401
    login(client)
    assert client.put('/api/state',json={'state':sample().model_dump(mode='json')},headers={'x-csrf-token':''}).status_code==403
    assert client.put('/api/state',json={'state':sample().model_dump(mode='json')}).status_code==200
    r=client.get('/api/state').json()
    assert r['state']['version']==1
    assert r['report']['totals']['net_profit']=='500'
    assert client.put('/api/state',json={'state':sample().model_dump(mode='json')}).status_code==409
    assert client.get('/assets/.env').status_code==404
    assert client.post('/api/logout').status_code==200
    assert client.get('/api/state').status_code==401

def test_import_preview_commit_duplicate(client,monkeypatch):
    import server
    async def fake(content,settings,hint):
        return {'inventory':[{'item':'Tea','quantity':10,'rmb_price':2,'exchange_rate':40,'pk_sell':100,'low_stock':5}],'bills':[],'payments':[],'accounts':[],'expenses':[],'warnings':[]}
    monkeypatch.setattr(server,'extract',fake)
    login(client)
    draft=client.post('/api/extract',files={'file':('stock.csv',b'Item,Quantity,RMB Price,Exchange Rate,PK Sell\nTea,10,2,40,100','text/csv')})
    assert draft.status_code==200,draft.text
    d=draft.json()
    assert client.get('/api/state').json()['state']['inventory']==[]
    body={'version':d['version'],'source':d['source'],'hash':d['hash'],'rows':{k:d[k] for k in ['inventory','bills','payments','accounts','expenses']}}
    assert client.post('/api/import',json=body).status_code==200
    assert client.post('/api/import',json=body).status_code==409
    assert len(client.get('/api/state').json()['state']['inventory'])==1

def test_invalid_import_is_atomic(client):
    login(client)
    body={'version':0,'source':'bad','hash':'f'*64,'rows':{'inventory':[{'item':'Valid','quantity':1,'rmb_price':1,'exchange_rate':40,'pk_sell':100},{'item':'Invalid','quantity':None}]}}
    assert client.post('/api/import',json=body).status_code==422
    assert client.get('/api/state').json()['state']['inventory']==[]

def test_image_excel_parser_and_export(client):
    from PIL import Image
    import openpyxl
    im=io.BytesIO();Image.new('RGB',(100,100),'white').save(im,format='PNG')
    assert read_upload('shot.PNG',im.getvalue())['type']=='image_url'
    wb=openpyxl.Workbook();ws=wb.active;ws.title='Inventory List';ws.append(['Item','Quantity']);ws.append(['Tea',10]);buf=io.BytesIO();wb.save(buf)
    assert 'Tea' in read_upload('stock.xlsx',buf.getvalue())['text']
    login(client)
    assert client.put('/api/state',json={'state':sample().model_dump(mode='json')}).status_code==200
    r=client.get('/api/export');assert r.status_code==200
    out=openpyxl.load_workbook(io.BytesIO(r.content));assert 'ledger' in out.sheetnames
    assert out['inventory']['B2'].value=='Shirt'
    assert client.get('/api/export?format=json').json()['version']==1

def test_duplicate_bill_and_orphan_rejected():
    s=sample();s.bills.append(s.bills[0].model_copy(update={'id':'other'}))
    with pytest.raises(ValueError,match='Duplicate bill'):normalize(s)
    s=sample();s.inventory=[]
    with pytest.raises(ValueError,match='missing inventory'):normalize(s)

def test_redis_cas_command(monkeypatch):
    import asyncio
    from storage import Store
    monkeypatch.setenv('UPSTASH_REDIS_REST_URL','https://test.upstash.io')
    monkeypatch.setenv('UPSTASH_REDIS_REST_TOKEN','fake')
    store=Store();calls=[]
    async def fake(*args):calls.append(args);return 1
    store.command=fake
    asyncio.run(store.save(State(),0))
    assert calls[0][0]=='EVAL'
    assert json.loads(calls[0][-1])['version']==1

def test_vercel_requires_redis(monkeypatch):
    from storage import Store
    from fastapi import HTTPException
    monkeypatch.setenv('VERCEL','1');monkeypatch.delenv('UPSTASH_REDIS_REST_URL',raising=False);monkeypatch.delenv('UPSTASH_REDIS_REST_TOKEN',raising=False)
    with pytest.raises(HTTPException):Store()

def test_workbook_without_dimensions_reads_all_shuffled_sheets(client,monkeypatch):
    """Valid third-party XLSX may omit dimension; all three tables must survive."""
    import re, zipfile, openpyxl, server
    wb=openpyxl.Workbook()
    ws=wb.active;ws.title='Payment List'
    ws.append(['Date','Name','Details','Amount']);ws.append(['2026-10-22','Waiz','Cash',5000])
    ws=wb.create_sheet('Inventory List')
    ws.append(['Item','Quantity','RMB Price','PKR per RMB','PK Sell']);ws.append(['Tea',10,2,39,100])
    ws=wb.create_sheet('Bill List')
    ws.append(['Date','Name','Bill No','Amount']);ws.append(['2026-10-22','Waiz','TEST-101',12000])
    original=io.BytesIO();wb.save(original);patched=io.BytesIO()
    with zipfile.ZipFile(original) as src, zipfile.ZipFile(patched,'w') as dst:
        for name in src.namelist():
            data=src.read(name)
            if re.match(r'xl/worksheets/sheet\d+\.xml$',name):
                data=re.sub(rb'<dimension\b[^>]*/>',b'',data)
            dst.writestr(name,data)
    raw=patched.getvalue()
    probe=openpyxl.load_workbook(io.BytesIO(raw),read_only=True)
    assert probe.worksheets[0].max_row is None
    assert probe.worksheets[0].max_column is None
    probe.close()
    content=read_upload('shuffled.xlsx',raw)
    sheets=json.loads(content['text'].split(':\n',1)[1])
    assert [s['sheet'] for s in sheets]==['Payment List','Inventory List','Bill List']
    assert [len(s['rows']) for s in sheets]==[2,2,2]
    async def fake(content,settings,hint):
        assert hint=='auto'
        assert all(name in content['text'] for name in ['Payment List','Inventory List','Bill List'])
        return {'inventory':[{'item':'Tea','quantity':10,'rmb_price':2,'exchange_rate':39,'pk_sell':100}],
                'bills':[{'date':'2026-10-22','name':'Waiz','bill_no':'TEST-101','amount':12000,'lines':[]}],
                'payments':[{'date':'2026-10-22','name':'Waiz','details':'Cash','amount':5000}],
                'accounts':[],'expenses':[],'warnings':[]}
    monkeypatch.setattr(server,'extract',fake)
    login(client)
    response=client.post('/api/extract',files={'file':('shuffled.xlsx',raw,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},data={'hint':'auto'})
    assert response.status_code==200,response.text
    d=response.json()
    body={'version':d['version'],'source':d['source'],'hash':d['hash'],'rows':{k:d[k] for k in ['inventory','bills','payments','accounts','expenses']}}
    assert client.post('/api/import',json=body).status_code==200
    result=client.get('/api/state').json()
    assert all(len(result['state'][k])==1 for k in ['inventory','bills','payments'])
    assert result['report']['totals']['receivables']=='7000'


def test_native_excel_maps_shuffled_columns_without_ai(client, monkeypatch):
    import openpyxl, server
    async def forbidden(*args, **kwargs):
        raise AssertionError('Standard Excel headers must not require AI')
    monkeypatch.setattr(server,'extract',forbidden)
    import ai
    monkeypatch.setattr(ai,'extract',forbidden)
    wb=openpyxl.Workbook(); ws=wb.active;ws.title='Payment List'
    ws.append(['Details','Amount (PKR)','Name','Date'])
    ws.append(['Cash','Rs. 5,000','Waiz','2026-10-18'])
    ws=wb.create_sheet('Inventory List')
    ws.append(['PK Sell','RMB Price','Item','Quantity','PKR per RMB'])
    ws.append([650,10,'Cotton Shirt',20,39])
    ws=wb.create_sheet('Bill List')
    ws.append(['Bill No','Amount PKR','Customer','Date'])
    ws.append(['EX-101',12000,'Waiz','2026-10-18'])
    raw=io.BytesIO();wb.save(raw); login(client)
    d=client.post('/api/extract',files={'file':('all.xlsx',raw.getvalue())},data={'hint':'auto'})
    assert d.status_code==200,d.text
    preview=d.json()
    assert len(preview['sheets'])==3
    assert [s['entries'] for s in preview['sheets']]==[1,1,1]
    assert preview['inventory'][0]['item']=='Cotton Shirt'
    assert preview['payments'][0]['amount']=='5000'
    assert preview['bills'][0]['bill_no']=='EX-101'
    body={'version':preview['version'],'source':preview['source'],'hash':preview['hash'],
          'rows':{k:preview[k] for k in ['inventory','bills','payments','accounts','expenses']}}
    assert client.post('/api/import',json=body).status_code==200
    assert client.get('/api/state').json()['report']['totals']['receivables']=='7000'


def test_incorrect_dimensions_and_formula_cache_are_not_silent():
    import re, zipfile, openpyxl, asyncio
    from models import Settings
    from excel_import import extract_workbook
    wb=openpyxl.Workbook();ws=wb.active;ws.title='Inventory'
    ws.append(['Item','Quantity','RMB Price','Exchange Rate','PK Sell'])
    ws.append(['Tea',10,2,39,'=39*2+20'])
    ws.append(['Bag',5,3,39,200])
    raw=io.BytesIO();wb.save(raw);patch=io.BytesIO()
    with zipfile.ZipFile(raw) as src,zipfile.ZipFile(patch,'w') as dst:
        for n in src.namelist():
            data=src.read(n)
            if n=='xl/worksheets/sheet1.xml':
                data=re.sub(rb'<dimension\b[^>]*/>',b'<dimension ref="A1:A1"/>',data)
            dst.writestr(n,data)
    result=asyncio.run(extract_workbook('wrong-dimension.xlsx',patch.getvalue(),Settings()))
    assert len(result['inventory'])==2
    assert result['inventory'][0]['pk_sell'] is None
    assert any('Uncalculated formula' in w for w in result['warnings'])
    assert result['inventory'][1]['item']=='Bag'


def test_multiple_tables_in_one_sheet_and_derived_accounts():
    import openpyxl, asyncio
    from models import Settings
    from excel_import import extract_workbook
    wb=openpyxl.Workbook();ws=wb.active;ws.title='Register'
    for row in [['Payments'],['Name','Date','Details','Amount'],['Waiz','2026-10-22','Cash',5000],
                ['Inventory'],['Item','Qty','RMB Price','PKR/RMB','PK Sell'],['Tea',10,2,39,100],
                ['Bills'],['Bill No','Amount PKR','Customer','Date'],['MIX-1',12000,'Waiz','2026-10-22']]:ws.append(row)
    ws=wb.create_sheet('Accounts List');ws.append(['Name','Sheet no','Balance']);ws.append(['Waiz','01',7000])
    buf=io.BytesIO();wb.save(buf)
    result=asyncio.run(extract_workbook('mixed.xlsx',buf.getvalue(),Settings()))
    assert all(len(result[t])==1 for t in ['inventory','bills','payments'])
    assert result['accounts']==[]
    assert any('double-counting' in w for w in result['warnings'])


def test_larger_workbook_imports_over_120_records(client):
    import openpyxl
    wb=openpyxl.Workbook();ws=wb.active;ws.title='Inventory'
    ws.append(['Item','Quantity','RMB Price','Exchange Rate','PK Sell'])
    for i in range(130):ws.append([f'Item {i}',10,2,39,100])
    raw=io.BytesIO();wb.save(raw);login(client)
    response=client.post('/api/extract',files={'file':('large.xlsx',raw.getvalue())})
    assert response.status_code==200,response.text
    d=response.json();assert len(d['inventory'])==130
    body={'version':d['version'],'source':d['source'],'hash':d['hash'],
          'rows':{k:d[k] for k in ['inventory','bills','payments','accounts','expenses']}}
    r=client.post('/api/import',json=body)
    assert r.status_code==200,r.text
    assert r.json()['rows']==130


def test_bundled_multisheet_excel_end_to_end_without_groq(client, monkeypatch):
    import ai
    async def no_groq(*args, **kwargs):
        raise AssertionError('Bundled standard workbook must import without Groq')
    monkeypatch.setattr(ai,'completion',no_groq)
    path=Path(__file__).resolve().parents[1]/'examples/Sample_All_Sheets.xlsx'
    login(client)
    response=client.post('/api/extract',files={'file':(path.name,path.read_bytes())},data={'hint':'auto'})
    assert response.status_code==200,response.text
    result=response.json()
    assert [s['sheet'] for s in result['sheets']]==['Payment List','Inventory List','Bill List']
    assert [s['entries'] for s in result['sheets']]==[2,6,2]
    assert all(s['mapping']=='column headers' for s in result['sheets'])
    body={'version':result['version'],'source':result['source'],'hash':result['hash'],
          'rows':{k:result[k] for k in ['inventory','bills','payments','accounts','expenses']}}
    response=client.post('/api/import',json=body)
    assert response.status_code==200,response.text
    saved=client.get('/api/state').json()
    assert saved['report']['totals']['revenue']=='18000'
    assert saved['report']['totals']['payments']=='6000'
    assert saved['report']['totals']['receivables']=='12000'
    assert len(saved['state']['inventory'])==6
